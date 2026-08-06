import asyncio
import logging
from pathlib import Path

import openpyxl

from Fun.BaseTools.AsyncHTTP import AsyncHTTPManage

from .auth import AccessTokenManager
from .baidu_ocr import GeneralOCR
from .config import EditConfig, OCRPostConfig
from .excel_utils import find_data_start_row
from .exception import APIError, NetworkError
from .extractor import ExcelFloatImageExtractor

logger = logging.getLogger(__name__)


class BaiduOCRFacade:
    # 类变量：全局唯一的信号量，控制最大并发请求数（设为 2 以匹配百度免费 QPS）
    _ocr_semaphore = asyncio.Semaphore(2)  
    def __init__(self,edit_config:EditConfig,http_manager:AsyncHTTPManage):
        self._http_manager = http_manager
        self._post_prototype  = OCRPostConfig()
        self._edit_config = edit_config
        self._token_manager = AccessTokenManager(http_manager)

    async def _recognize_with_limit(self,img_base64:bytes)->dict:
        """内部方法：使用信号量包裹OCR请求"""
        async with self._ocr_semaphore:
            ocr = GeneralOCR(
                img_base64=img_base64,
                access_token_manager=self._token_manager,
                post_config=self._post_prototype,
                http_manager=self._http_manager,
                edit_config=self._edit_config
            )
            # 注意：这里直接返回识别结果，信号量会在 with 块结束后自动释放
            return await ocr.recognize()
    async def process_excel(self,file_path:str,sheet_name:str)->Path:
        with ExcelFloatImageExtractor(file_path) as extractor:
            images_info  = extractor.get_single_sheet_floating_images(sheet_name)
            image_data_list = []
            for img_info in images_info:
                row = img_info["row"]
                img_base64 = extractor.get_encoded_image_data(img_info["path"])
                if img_base64:
                    image_data_list.append((row, img_base64))
                else:
                    logger.warning(f"第 {row} 行图片数据为空，跳过")

            
        wb = openpyxl.load_workbook(file_path)
        sheet = wb[sheet_name]
        
        data_start_row = find_data_start_row(sheet)
        if data_start_row is None:
            raise ValueError("未找到有效数据行，请确保 A 列有数字编号。")
        header_row = data_start_row - 1

        max_col = sheet.max_column
        start_col = max_col + 1
        field_names = self._edit_config.get_all_display_fields()

        for idx,field in enumerate(field_names):
            sheet.cell(row=header_row,column = start_col+idx,value=field)
            


        # 3. ★ 核心并发点：创建所有图片的异步任务，一次性提交给事件循环
        tasks = []
        for row, img_base64 in image_data_list:
            task = asyncio.create_task(self._recognize_with_limit(img_base64))
            tasks.append((row, task))  # 保存行号和任务

        # 4. 等待所有任务完成（并发执行，但受 Semaphore(2) 限制）
        results_by_row = {}
        for row, task in tasks:
            try:
                result = await task
                results_by_row[row] = result
            except (NetworkError, APIError) as e:
                logger.error(f"第 {row} 行识别失败: {e}")
        field_to_col = {field: start_col + idx for idx, field in enumerate(field_names)}
        for row, results in results_by_row.items():
            type, data = results
            if type == 'idcard':
                mapping = self._edit_config.get_idcard_mapping()   # [(display_name, api_key), ...]
                for display_name, _ in mapping:
                    col = field_to_col.get(display_name)
                    if col is not None:
                        sheet.cell(row=row, column=col, value=data.get(display_name, ''))
            elif type == 'bankcard':
                mapping = self._edit_config.get_bankcard_mapping()
                for display_name, _ in mapping:
                    col = field_to_col.get(display_name)
                    if col is not None:
                        sheet.cell(row=row, column=col, value=data.get(display_name, ''))
        # 5. 保存文件
        output_path = Path(file_path).parent / f"{Path(file_path).stem}_processed.xlsx"
        wb.save(output_path)
        return output_path    